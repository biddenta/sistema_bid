"""
Exportador de Matches para Excel
Gera arquivos Excel formatados com dados de matching e estatísticas
"""

import sys
from pathlib import Path

# Adiciona o diretório raiz ao path
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))

import sqlite3
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def exportar_matches_excel():
    """Exporta tabela de matches para Excel formatado"""
    
    conn = sqlite3.connect('match_crew.db')
    
    print("\n" + "=" * 80)
    print("📊 EXPORTANDO MATCHES PARA EXCEL")
    print("=" * 80 + "\n")
    
    # Buscar todos os matches
    print("🔍 Carregando matches...")
    
    query_matches = """
        SELECT 
            pm.id_match as 'ID Match',
            pm.nome_produto as 'Nome do Produto',
            pm.categoria as 'Categoria',
            pm.subcategoria as 'Subcategoria',
            pm.marca as 'Marca',
            pm.embalagem as 'Embalagem',
            pm.total_sites as 'Nº Sites',
            pm.total_produtos as 'Nº Produtos',
            ROUND(pm.score_match, 2) as 'Score',
            pm.url_cremer as 'URL Dental Cremer',
            pm.url_speed as 'URL Dental Speed',
            pm.url_medsul as 'URL Dental Medsul',
            pm.url_proclin as 'URL Dental Proclin',
            pm.url_dentalshop as 'URL Dental Shop',
            pm.url_apoiodental as 'URL Apoio Dental',
            pm.url_interdental as 'URL Interdental',
            pm.url_surya as 'URL Surya',
            pm.metodo_matching as 'Método',
            DATE(pm.data_criacao) as 'Data Criação'
        FROM produtos_mestre pm
        ORDER BY pm.total_produtos DESC, pm.total_sites DESC
    """
    
    df_matches = pd.read_sql_query(query_matches, conn)
    
    print(f"   ✅ {len(df_matches):,} matches carregados\n")
    
    # Buscar preços por match
    print("💰 Calculando preços...")
    
    query_precos = """
        SELECT 
            pm.id_match,
            p.site,
            COALESCE(p.preco_promocional, p.preco, 0) as preco
        FROM produtos_mestre pm
        JOIN produtos p ON p.id_match = pm.id_match
        WHERE p.tem_matching = 1 AND COALESCE(p.preco_promocional, p.preco, 0) > 0
        ORDER BY pm.id_match
    """
    
    df_precos = pd.read_sql_query(query_precos, conn)
    
    # Calcular preços min/max por match
    precos_agregados = df_precos.groupby('id_match').agg({
        'preco': ['min', 'max']
    }).reset_index()
    
    precos_agregados.columns = ['ID Match', 'Preço Mínimo', 'Preço Máximo']
    
    # Calcular variação percentual
    precos_agregados['Variação (%)'] = (
        (precos_agregados['Preço Máximo'] - precos_agregados['Preço Mínimo']) / 
        precos_agregados['Preço Mínimo'] * 100
    ).round(1)
    
    # Encontrar melhor site (menor preço)
    idx_min = df_precos.groupby('id_match')['preco'].idxmin()
    melhor_site = df_precos.loc[idx_min, ['id_match', 'site']]
    melhor_site.columns = ['ID Match', 'Melhor Site']
    
    # Formatar preços
    precos_agregados['Preço Mínimo'] = precos_agregados['Preço Mínimo'].apply(
        lambda x: f"R$ {x:.2f}" if pd.notna(x) else ""
    )
    precos_agregados['Preço Máximo'] = precos_agregados['Preço Máximo'].apply(
        lambda x: f"R$ {x:.2f}" if pd.notna(x) else ""
    )
    precos_agregados['Variação (%)'] = precos_agregados['Variação (%)'].apply(
        lambda x: f"{x:.1f}%" if pd.notna(x) and x > 0 else ""
    )
    
    print(f"   ✅ Preços calculados\n")
    
    # Merge dos dataframes
    df_final = df_matches.merge(precos_agregados, on='ID Match', how='left')
    df_final = df_final.merge(melhor_site, on='ID Match', how='left')
    
    # Reordenar colunas
    cols_ordem = [
        'ID Match', 'Nome do Produto', 'Categoria', 'Subcategoria', 'Marca', 'Embalagem',
        'Nº Sites', 'Nº Produtos', 'Score',
        'Preço Mínimo', 'Preço Máximo', 'Variação (%)', 'Melhor Site',
        'URL Dental Cremer', 'URL Dental Speed', 'URL Dental Medsul', 'URL Dental Proclin',
        'URL Dental Shop', 'URL Apoio Dental', 'URL Interdental', 'URL Surya',
        'Método', 'Data Criação'
    ]
    
    df_final = df_final[cols_ordem]
    
    # Criar estatísticas
    print("📊 Gerando estatísticas...")
    
    cursor = conn.cursor()
    
    cursor.execute("SELECT COUNT(*) FROM produtos")
    total_produtos = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM produtos WHERE tem_matching = 1")
    com_matching = cursor.fetchone()[0]
    
    # Distribuição por sites
    cursor.execute("""
        SELECT total_sites, COUNT(*) as total
        FROM produtos_mestre
        GROUP BY total_sites
        ORDER BY total_sites DESC
    """)
    dist_sites = cursor.fetchall()
    
    # Top categorias
    cursor.execute("""
        SELECT categoria, COUNT(*) as total
        FROM produtos_mestre
        WHERE categoria IS NOT NULL
        GROUP BY categoria
        ORDER BY total DESC
        LIMIT 10
    """)
    top_categorias = cursor.fetchall()
    
    # Criar DataFrame de estatísticas
    stats_data = []
    stats_data.append(['ESTATÍSTICAS GERAIS', ''])
    stats_data.append(['Total de Produtos', f"{total_produtos:,}"])
    stats_data.append(['Produtos com Matching', f"{com_matching:,}"])
    stats_data.append(['Taxa de Matching', f"{com_matching/total_produtos*100:.2f}%"])
    stats_data.append(['Total de Grupos', f"{len(df_matches):,}"])
    stats_data.append(['', ''])
    stats_data.append(['DISTRIBUIÇÃO POR Nº DE SITES', ''])
    
    for sites, count in dist_sites:
        stats_data.append([f"{sites} sites", f"{count:,} grupos"])
    
    stats_data.append(['', ''])
    stats_data.append(['TOP 10 CATEGORIAS', ''])
    
    for cat, count in top_categorias:
        stats_data.append([cat or 'Sem categoria', f"{count:,} grupos"])
    
    df_stats = pd.DataFrame(stats_data, columns=['Descrição', 'Valor'])
    
    print(f"   ✅ Estatísticas geradas\n")
    
    # Salvar Excel
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"exports/matches_produtos_{timestamp}.xlsx"
    
    print(f"💾 Salvando arquivo: {filename}")
    
    # Criar diretório exports se não existir
    import os
    os.makedirs('exports', exist_ok=True)
    
    # Exportar para Excel
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        df_final.to_excel(writer, sheet_name='Matches de Produtos', index=False)
        df_stats.to_excel(writer, sheet_name='Estatísticas', index=False)
    
    # Aplicar formatação
    print("🎨 Aplicando formatação...")
    
    wb = load_workbook(filename)
    ws = wb['Matches de Produtos']
    
    # Estilos
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Formatar cabeçalhos
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Ajustar larguras
    larguras = {
        'A': 15, 'B': 50, 'C': 25, 'D': 25, 'E': 20, 'F': 15,
        'G': 10, 'H': 12, 'I': 8, 'J': 15, 'K': 15, 'L': 12, 'M': 20,
        'N': 15, 'O': 15, 'P': 15, 'Q': 15, 'R': 15, 'S': 15, 'T': 15, 'U': 15,
        'V': 20, 'W': 12
    }
    
    for col, width in larguras.items():
        ws.column_dimensions[col].width = width
    
    # Congelar primeira linha
    ws.freeze_panes = 'A2'
    
    # Adicionar filtros
    ws.auto_filter.ref = ws.dimensions
    
    # Formatar estatísticas
    ws_stats = wb['Estatísticas']
    
    for row in ws_stats.iter_rows(min_row=1):
        if 'ESTATÍSTICAS' in str(row[0].value) or 'DISTRIBUIÇÃO' in str(row[0].value) or 'TOP' in str(row[0].value):
            row[0].font = Font(bold=True, size=12)
            row[0].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    ws_stats.column_dimensions['A'].width = 30
    ws_stats.column_dimensions['B'].width = 20
    
    wb.save(filename)
    
    print(f"   ✅ Formatação aplicada\n")
    
    print(f"{'='*80}")
    print("✅ EXPORTAÇÃO CONCLUÍDA COM SUCESSO!")
    print(f"{'='*80}\n")
    print(f"   📁 Arquivo: {filename}")
    print(f"   📊 Matches exportados: {len(df_matches):,}")
    print(f"   📄 Planilhas: 2 (Matches + Estatísticas)")
    print(f"   💾 Linhas totais: {len(df_final) + len(df_stats)}")
    print(f"\n{'='*80}\n")
    
    conn.close()


if __name__ == "__main__":
    try:
        exportar_matches_excel()
    except Exception as e:
        print(f"\n❌ Erro: {e}")
        import traceback
        traceback.print_exc()
